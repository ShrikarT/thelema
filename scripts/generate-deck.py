import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    BG_DARK = RGBColor(11, 15, 25)        # #0B0F19
    CARD_BG = RGBColor(19, 27, 46)        # #131B2E
    CARD_BORDER = RGBColor(30, 41, 59)    # #1E293B
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8
    TEXT_BODY = RGBColor(203, 213, 225)   # #CBD5E1
    ACCENT_BLUE = RGBColor(56, 189, 248)  # #38BDF8
    ACCENT_GREEN = RGBColor(16, 185, 129) # #10B981
    ACCENT_PURPLE = RGBColor(129, 140, 248)

    evidence_dir = os.path.abspath("docs/evidence")

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    # =========================================================================
    # SLIDE 1: Title
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    # Decorative glow card
    add_card(s1, Inches(1.2), Inches(1.2), Inches(10.933), Inches(5.1), RGBColor(15, 22, 38))

    tb = s1.shapes.add_textbox(Inches(1.8), Inches(1.6), Inches(9.733), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "ETHONLINE 2026 SUBMISSION"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE

    p1 = tf.add_paragraph()
    p1.text = "THELEMA"
    p1.font.size = Pt(56)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = "Trade the impact. Not just the odds."
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_GREEN
    p2.space_after = Pt(16)

    p3 = tf.add_paragraph()
    p3.text = "Conditional synthetic prediction markets deployed on Arc Testnet, powered by The Graph decentralized gateway consensus and Chainlink CRE confidential workflows."
    p3.font.size = Pt(15)
    p3.font.color.rgb = TEXT_BODY
    p3.space_after = Pt(28)

    p4 = tf.add_paragraph()
    p4.text = "Arc Testnet (Chain ID 5042002)   |   The Graph (Messari Schema)   |   Chainlink CRE (WASM)   |   100% Solvency Vaults"
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 2: The Problem
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)

    # Header
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "THE PROBLEM"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p = tf.add_paragraph()
    p.text = "Prediction Markets Predict Odds, Not Consequences"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    # Left Card - Binary Today
    add_card(s2, Inches(1.0), Inches(2.2), Inches(5.4), Inches(4.5))
    tb_left = s2.shapes.add_textbox(Inches(1.3), Inches(2.4), Inches(4.8), Inches(4.0))
    tf_left = tb_left.text_frame
    tf_left.word_wrap = True

    p = tf_left.paragraphs[0]
    p.text = "Binary Markets Today (e.g. Polymarket)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(248, 113, 113) # Red accent
    p.space_after = Pt(18)

    bullets_left = [
        ("Binary $1 / $0 Payout:", "Tells you if policy passes, but zero information on economic impact."),
        ("No Asset Pricing:", "Cannot price how Nvidia, semiconductor equities, or commodities react to the decision."),
        ("Speculative Only:", "Investors cannot use binary outcome tokens to hedge physical portfolio exposure.")
    ]
    for title, desc in bullets_left:
        p = tf_left.add_paragraph()
        p.text = f"• {title} "
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(4)
        
        p_desc = tf_left.add_paragraph()
        p_desc.text = f"   {desc}"
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = TEXT_MUTED
        p_desc.space_after = Pt(12)

    # Right Card - Thelema Solution
    add_card(s2, Inches(6.9), Inches(2.2), Inches(5.4), Inches(4.5))
    tb_right = s2.shapes.add_textbox(Inches(7.2), Inches(2.4), Inches(4.8), Inches(4.0))
    tf_right = tb_right.text_frame
    tf_right.word_wrap = True

    p = tf_right.paragraphs[0]
    p.text = "What Capital Actually Needs (THELEMA)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(18)

    bullets_right = [
        ("Conditional Synthetic Assets:", "Trade outcome-contingent shares: sNVDA if policy passes vs sNVDA if policy fails."),
        ("Quantified Price Impact:", "Derives Δ Impact = E[S | YES] − E[S | NO] directly from continuous AMM pool reserves."),
        ("True Portfolio Hedging:", "Holders of physical stocks or treasuries can hedge regulatory risk with mathematical certainty.")
    ]
    for title, desc in bullets_right:
        p = tf_right.add_paragraph()
        p.text = f"• {title} "
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(4)
        
        p_desc = tf_right.add_paragraph()
        p_desc.text = f"   {desc}"
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = TEXT_BODY
        p_desc.space_after = Pt(12)

    # =========================================================================
    # SLIDE 3: The Pricing Identity & Mathematics
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)

    tb = s3.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "THE PRICING MECHANISM"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p = tf.add_paragraph()
    p.text = "Same Asset. Different Futures. Four Verified Numbers."
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    # 4 Stat Cards
    stat_cards = [
        ("Event Probability", "54.7%", "P(event)", "Binary YES pool midprice", ACCENT_BLUE),
        ("Asset if YES", "$212.92", "E[S | YES]", "Conditional price if allowed", ACCENT_GREEN),
        ("Asset if NO", "$96.46", "E[S | NO]", "Conditional price if restricted", RGBColor(251, 146, 60)),
        ("Conditional Impact", "+$116.47", "Δ Impact", "YES world minus NO world", RGBColor(168, 85, 247))
    ]
    card_w = Inches(2.65)
    for i, (title, val, notation, sub, col) in enumerate(stat_cards):
        cx = Inches(1.0 + i * 2.9)
        add_card(s3, cx, Inches(2.2), card_w, Inches(2.5))
        
        tb_c = s3.shapes.add_textbox(cx + Inches(0.2), Inches(2.3), card_w - Inches(0.4), Inches(2.3))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        
        p = tf_c.paragraphs[0]
        p.text = title.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED
        
        p = tf_c.add_paragraph()
        p.text = val
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(4)
        
        p = tf_c.add_paragraph()
        p.text = notation
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        
        p = tf_c.add_paragraph()
        p.text = sub
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED

    # Bottom Mathematical Invariant Card
    add_card(s3, Inches(1.0), Inches(5.0), Inches(11.333), Inches(1.8))
    tb_bottom = s3.shapes.add_textbox(Inches(1.3), Inches(5.1), Inches(10.733), Inches(1.6))
    tf_b = tb_bottom.text_frame
    tf_b.word_wrap = True

    p = tf_b.paragraphs[0]
    p.text = "The Pricing Identity & Solvency Guarantee:"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p = tf_b.add_paragraph()
    p.text = "P × E[S | YES]  +  (1 − P) × E[S | NO]  =  Implied Spot Index ($160.21)"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_after = Pt(6)

    p = tf_b.add_paragraph()
    p.text = "Mathematical Vault Solvency: Depositing collateral mints complete sets: Y + N + R capped at $500. Winning asset shares pay min(index, cap); losing pays 0; residual R pays max(0, cap − min(index, cap)). Collateral is 100% conserved."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_BODY

    # =========================================================================
    # SLIDE 4: Sponsor Architecture
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)

    tb = s4.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "INFRASTRUCTURE & SPONSOR TRACKS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p = tf.add_paragraph()
    p.text = "Three Rigorous Integrations Built from Scratch"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    col_w = Inches(3.6)
    tracks = [
        ("Arc Testnet", "Best Stablecoin-Native DeFi Pool", ACCENT_BLUE, [
            "6 Smart Contracts Active: BinaryAMM, ShareAMM, Vaults, and DemoOracle.",
            "Native USDC Precompile (0x3600...): Direct 6-decimal collateral and native gas.",
            "Live On-Chain Trade: 1.00 USDC demo trade executed & confirmed on Arcscan in block 61728406.",
            "Zero Mocked State: Verified bytecode, balances, and nonces on Chain ID 5042002."
        ]),
        ("The Graph", "Messari Standardized Schema", ACCENT_GREEN, [
            "Cross-Protocol Engine: Queries Uniswap V3 and SushiSwap pools on Arbitrum One.",
            "Standardized Messari DEX Schema: Normalized fields across different DEX protocols.",
            "Real-Time Spread Consensus: Verifies pools agree within 0.25% spread (threshold 1.5%).",
            "Decentralized Gateway: Server-side queries keep production keys out of client bundle."
        ]),
        ("Chainlink CRE", "Confidential Workflows", ACCENT_PURPLE, [
            "Confidential Size Clipping: Hides order notional inside TEE to prevent front-running.",
            "WASM Policy Compilation: Pure TypeScript policy compiled to standalone WebAssembly.",
            "Official CLI Simulation: Tested via Chainlink CRE CLI simulator (halted at network boundary).",
            "Honest Hackathon Reporting: Production contract standby pending organization deploy access."
        ])
    ]

    for i, (name, track_name, col, items) in enumerate(tracks):
        cx = Inches(1.0 + i * 3.866)
        add_card(s4, cx, Inches(2.2), col_w, Inches(4.6))
        
        tb_t = s4.shapes.add_textbox(cx + Inches(0.2), Inches(2.4), col_w - Inches(0.4), Inches(4.2))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        
        p = tf_t.paragraphs[0]
        p.text = name
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = col
        
        p = tf_t.add_paragraph()
        p.text = track_name
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(14)
        
        for item in items:
            p = tf_t.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_BODY
            p.space_after = Pt(8)

    # =========================================================================
    # SLIDE 5: Live Market Walkthrough
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)

    tb = s5.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.333), Inches(1.0))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "LIVE PRODUCT DEMONSTRATION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p = tf.add_paragraph()
    p.text = "Active On-Chain Prediction Market on Arc Testnet"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    # Left content box
    add_card(s5, Inches(1.0), Inches(1.8), Inches(4.8), Inches(5.1))
    tb_walk = s5.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(4.4), Inches(4.7))
    tf_w = tb_walk.text_frame
    tf_w.word_wrap = True

    p = tf_w.paragraphs[0]
    p.text = "Market Parameters & Features"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(12)

    steps = [
        ("Target Market:", "Will the US allow advanced AI chip sales to China by 31 Dec 2026?"),
        ("Active Trading Window:", "Event deadline: 31 Dec 2026 | Trading cutoff: 1 Jan 2027."),
        ("Real Web3 Wallet Integration:", "Connects MetaMask on Chain ID 5042002. Displays live on-chain token balances."),
        ("Automated Quotation:", "Calculates precise 30 bps LP fee, price impact, and minimum shares received."),
        ("One-Click Execution:", "Direct smart contract interactions without custodial intermediaries.")
    ]
    for h, b in steps:
        p = tf_w.add_paragraph()
        p.text = f"{h} "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(2)
        
        p2 = tf_w.add_paragraph()
        p2.text = f"{b}"
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_BODY
        p2.space_after = Pt(8)

    # Right Image
    img_path1 = os.path.join(evidence_dir, "submission-01-arc-market.png")
    if os.path.exists(img_path1):
        s5.shapes.add_picture(img_path1, Inches(6.1), Inches(1.8), width=Inches(6.2))

    # =========================================================================
    # SLIDE 6: On-Chain Proof & Reconciliation
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)

    tb = s6.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.333), Inches(1.0))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "ON-CHAIN PROOF & AUDIT"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p = tf.add_paragraph()
    p.text = "Verified Demonstration Trade & Spending Reconciliation"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    # Left content box
    add_card(s6, Inches(1.0), Inches(1.8), Inches(4.8), Inches(5.1))
    tb_proof = s6.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(4.4), Inches(4.7))
    tf_p = tb_proof.text_frame
    tf_p.word_wrap = True

    p = tf_p.paragraphs[0]
    p.text = "Trade Proof & Holdings"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(12)

    proofs = [
        ("Trade Hash:", "0xe4208cc844dbfec21d18e38c20630388c2a7c7083bb17583fbe6dfc5a55190a6"),
        ("Mined Block:", "Block 61728406 on Arc Testnet (Success)."),
        ("Swap Execution:", "Deposited 1.000000 USDC -> Received 1.903611 tYES tokens."),
        ("Verified On-Chain Holdings:", "11.9036 tYES, 10.0000 tNO, 0.1000 R."),
        ("Audited Spending Ceiling:", "Total spent: 173.707142 USDC <= 180.000000 USDC ceiling. Balance: 22.292857 USDC.")
    ]
    for h, b in proofs:
        p = tf_p.add_paragraph()
        p.text = f"{h} "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(2)
        
        p2 = tf_p.add_paragraph()
        p2.text = f"{b}"
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_BODY
        p2.space_after = Pt(8)

    # Right Image
    img_path4 = os.path.join(evidence_dir, "submission-04-onchain-proof.png")
    if os.path.exists(img_path4):
        s6.shapes.add_picture(img_path4, Inches(6.1), Inches(1.8), width=Inches(6.2))

    # =========================================================================
    # SLIDE 7: Summary & Conclusion
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)

    tb = s7.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "CONCLUSION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p = tf.add_paragraph()
    p.text = "The Next Generation of Prediction Markets"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    # 3 Feature Cards
    summary_cards = [
        ("100% Verified Quality", ACCENT_BLUE, [
            "238 core unit & integration tests passing.",
            "24 Foundry Solidity invariant tests passing.",
            "18 Playwright end-to-end browser tests passing.",
            "Green GitHub Actions CI on every commit."
        ]),
        ("Live & Deployed", ACCENT_GREEN, [
            "Active on-chain contracts on Arc Testnet.",
            "Native 6-decimal USDC collateral & gas precompile.",
            "Live hosted application on Render.",
            "Real-time The Graph decentralized reference feeds."
        ]),
        ("Future Roadmap", ACCENT_PURPLE, [
            "Deploy live CRE TEE workflows once org access opens.",
            "Expand to macro tariffs, interest rates & chip governance.",
            "Permissionless creation of custom impact pairs.",
            "Institutional hedging integrations."
        ])
    ]

    for i, (title, col, bullets) in enumerate(summary_cards):
        cx = Inches(1.0 + i * 3.866)
        add_card(s7, cx, Inches(2.2), col_w, Inches(3.6))
        
        tb_sc = s7.shapes.add_textbox(cx + Inches(0.2), Inches(2.4), col_w - Inches(0.4), Inches(3.2))
        tf_sc = tb_sc.text_frame
        tf_sc.word_wrap = True
        
        p = tf_sc.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(12)
        
        for b in bullets:
            p = tf_sc.add_paragraph()
            p.text = f"• {b}"
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_BODY
            p.space_after = Pt(8)

    # Bottom Contact & Link Bar
    add_card(s7, Inches(1.0), Inches(6.0), Inches(11.333), Inches(0.9), RGBColor(15, 22, 38))
    tb_links = s7.shapes.add_textbox(Inches(1.3), Inches(6.15), Inches(10.733), Inches(0.6))
    tf_l = tb_links.text_frame
    p = tf_l.paragraphs[0]
    p.text = "Live Hosted Application:  thelema.onrender.com      |      GitHub Repository:  github.com/ShrikarT/thelema"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    out_file = "THELEMA-Presentation.pptx"
    prs.save(out_file)
    print(f"Presentation saved to {out_file}")

if __name__ == "__main__":
    build_deck()
